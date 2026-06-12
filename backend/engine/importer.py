"""Acquisition import pipeline (LoadData port) + rule-driven synthetic sample
generator that produces files in the exact TRIE format the original expects:
row 1 = family, row 2 = channel, keys are "Family, Channel" pairs.
"""
import io
import random
import openpyxl

from .classifier import classify_event


def parse_trie(file_bytes):
    """Parse a TRIE acquisition workbook -> (events list of channel dicts, headers)."""
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True, read_only=True)
    sheet_name = None
    for name in wb.sheetnames:
        if name.strip().upper() == "TRIE":
            sheet_name = name
            break
    if sheet_name is None:
        raise ValueError("The acquisition file has no 'TRIE' tab")
    ws = wb[sheet_name]
    rows = list(ws.iter_rows(values_only=True))
    if len(rows) < 3:
        raise ValueError("TRIE tab has no event rows")
    families, channels = rows[0], rows[1]
    headers = []
    for i in range(len(channels)):
        fam = str(families[i]).strip() if i < len(families) and families[i] is not None else ""
        chan = str(channels[i]).strip() if channels[i] is not None else ""
        if not chan:
            headers.append(None)
        else:
            headers.append(f"{fam}, {chan}" if fam else chan)
    events = []
    for row in rows[2:]:
        if row is None or all(v is None for v in row):
            continue
        rec = {}
        for i, h in enumerate(headers):
            if h is None or i >= len(row):
                continue
            v = row[i]
            if v is not None:
                rec[h.replace(".", "\u00b7")] = v
        if rec:
            events.append(rec)
    return events, [h for h in headers if h]


# ------------------------------------------------------- synthetic generator

WEAK_SDVS = {"COAST - BRAKE-ON DOWNSHIFT", "POWER-ON UPSHIFT COLD",
             "DRIVE AWAY CREEP ENG ON", "TIP IN AT DECELERATION",
             "CONVERTER RELEASE", "SHIFT ABORT", "MANEUVERING", "AUTO STOP"}
MID_SDVS = {"POWER-ON UPSHIFT", "POWER-ON DOWNSHIFT", "LEVER CHANGE",
            "DECEL CST BRAKE", "ACCEL LOAD INCREASE", "TIP OUT UPSHIFT",
            "CONVERTER LOCK UP", "DASS ENG ON", "DA ROLLING START",
            "LOAD REVERSAL DOWNSHIFT"}


def _channel_defaults(rules):
    """Neutral defaults for every channel referenced by any rule, so the
    missing-channel auto-pass ('FFF' sentinel) cannot mis-route events."""
    numeric = {}
    for rule in rules:
        for cond in rule.get("conditions", []):
            key = str(cond.get("column", "")).strip().replace(".", "\u00b7")
            op = cond.get("op", "")
            is_num = op in ("INFERIEUR A", "INFERIEUR OU EGAL A",
                            "SUPERIEUR A", "SUPERIEUR OU EGAL A")
            numeric[key] = numeric.get(key, True) and is_num
    return {k: (0.0 if is_num else "----") for k, is_num in numeric.items()}


def _candidates_for(conds):
    """Possible values for one channel given all its conditions."""
    out = []
    for cond in conds:
        op = cond.get("op", "EGAL A")
        val = cond.get("value", "")
        if op == "EGAL A":
            out.append(val)
        elif op == "CONTIENT":
            out.append(str(val).split(";")[0])
        try:
            f = float(str(val).replace(",", "."))
        except (TypeError, ValueError):
            continue
        if op == "SUPERIEUR A":
            out += [f + 10, f + 0.5, f + 30]
        elif op == "SUPERIEUR OU EGAL A":
            out += [f, f + 10]
        elif op == "INFERIEUR A":
            out += [f - 1, f - 0.3, f - 10]
        elif op == "INFERIEUR OU EGAL A":
            out += [f, f - 1]
    out += ["55", 0.0, "----", ""]
    seen, uniq = set(), []
    for v in out:
        k = repr(v)
        if k not in seen:
            seen.add(k)
            uniq.append(v)
    return uniq


def _channel_options(rule):
    """For each channel of a rule, the candidate values that satisfy all of
    the rule's groups touching that channel (group = OR, groups = AND)."""
    from .classifier import check_condition
    by_chan = {}
    groups = {}
    for cond in rule.get("conditions", []):
        groups.setdefault(cond.get("group", 1), []).append(cond)
        key = str(cond.get("column", "")).strip().replace(".", "\u00b7")
        by_chan.setdefault(key, []).append(cond)
    options = {}
    for key, conds in by_chan.items():
        # groups in which this channel appears
        chan_groups = [g for g in groups.values()
                       if any(str(c.get("column", "")).strip().replace(".", "\u00b7") == key for c in g)]
        good = []
        for cand in _candidates_for(conds):
            ok = True
            for g in chan_groups:
                # group satisfied if ANY of its conditions passes; conditions on
                # other channels are ignored here (handled by their own channel)
                mine = [c for c in g if str(c.get("column", "")).strip().replace(".", "\u00b7") == key]
                others = len(mine) < len(g)
                if not any(check_condition(cand, c.get("op", "EGAL A"), c.get("value", "")) for c in mine) \
                        and not others:
                    ok = False
                    break
            if ok:
                good.append(cand)
        options[key] = good or ["55"]
    return options


def _base_channels_for_rule(rule, defaults, rules, rng=None):
    """Construct channels that classify to this rule, trying candidate
    combinations until verification passes. Returns (channels, fixed) or None."""
    options = _channel_options(rule)
    keys = list(options.keys())
    rng = rng or random.Random(0)
    attempts = [[opts[0] for opts in options.values()]]
    for _ in range(40):
        attempts.append([rng.choice(options[k]) for k in keys])
    for combo in attempts:
        channels = dict(defaults)
        for k, v in zip(keys, combo):
            channels[k] = v
        if classify_event(channels, rules) == rule["sdv"]:
            return channels, set(keys)
    return None


def generate_sample_events(definitions, structure, canon_map, per_sdv=8, seed=7):
    """Generate classified-by-construction events for every active rule."""
    rng = random.Random(seed)
    rules = sorted(definitions, key=lambda r: r.get("order", 0))
    defaults = _channel_defaults(rules)
    all_events = []
    for rule in rules:
        if not rule.get("active", True):
            continue
        built = _base_channels_for_rule(rule, defaults, rules, rng)
        if built is None:
            continue  # an earlier rule steals it whatever the channel values
        base, fixed = built
        canon = canon_map.get(rule["sdv"].strip().upper(), rule["sdv"])
        spec = structure.get(canon, {"data": [], "criteria": []})
        if str(canon).strip().upper() in WEAK_SDVS:
            quality = 6.6
        elif str(canon).strip().upper() in MID_SDVS:
            quality = 7.6
        else:
            quality = 9.0
        n = rng.randint(max(3, per_sdv - 3), per_sdv + 4)
        for i in range(n):
            ev = dict(base)
            speed = round(rng.uniform(0, 130), 1)
            accel = round(rng.uniform(-3.2, 3.2), 2)
            g_old = rng.randint(1, 7)
            gen = {
                "gear old": g_old, "gear new": min(8, g_old + rng.choice([-1, 0, 1])),
                "Vehicle Speed": speed, "Throttle Position": round(rng.uniform(0, 100)),
                "Turbine Speed": int(900 + speed * 30), "Front shaft Torque": round(rng.uniform(-80, 350)),
                "Shift delay (sec)": round(rng.uniform(0.1, 0.9), 2),
                "Shift duration (sec)": round(rng.uniform(0.2, 1.2), 2),
                "Rpm flare": round(rng.uniform(0, 120)),
                "Gearbox Temperature": round(rng.uniform(55, 95)),
                "Engine Temperature": round(rng.uniform(60, 100)),
                "Learn Rule": rng.choice(["LR1", "LR2", "----"]),
                "Start time of the Sub Event": round(rng.uniform(0, 3600), 1),
                "Acquisition Name": "DEMO_RUN_%02d" % rng.randint(1, 4),
            }
            for d in spec.get("data", []):
                imp = d.get("import")
                if not imp:
                    continue
                key = imp.replace(".", "\u00b7")
                if key in fixed or key in ev:
                    continue
                ev[key] = gen.get(d["name"], round(rng.uniform(0, 100), 1))
            ev.setdefault("AccelerationChassis", accel)
            if "P: vehicle speed, vehicle Speed" not in fixed:
                ev["P: vehicle speed, vehicle Speed"] = speed
            for cr in spec.get("criteria", []):
                imp = cr.get("import") or (", " + cr["name"])
                chan = imp.split(",", 1)[1].strip() if "," in imp else imp
                key = chan.replace(".", "\u00b7")
                if key not in ev:
                    ev[key] = round(min(10.0, max(1.0, rng.gauss(quality, 1.15))), 1)
            ev.setdefault("Note, Event rating", round(min(10.0, max(1.0, rng.gauss(quality, 0.9))), 1))
            if classify_event(ev, rules) != rule["sdv"]:
                continue  # random fill broke the routing; drop this variant
            all_events.append(ev)
    return all_events


def write_sample_workbook(path, events):
    """Write events to an xlsx with a TRIE tab (family row + channel row)."""
    keys = []
    seen = set()
    for ev in events:
        for k in ev.keys():
            if k not in seen:
                seen.add(k)
                keys.append(k)
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "TRIE"
    fams, chans = [], []
    for k in keys:
        kk = k.replace("\u00b7", ".")
        if "," in kk:
            fam, chan = kk.split(",", 1)
            fams.append(fam.strip())
            chans.append(chan.strip())
        else:
            fams.append("")
            chans.append(kk)
    ws.append(fams)
    ws.append(chans)
    for ev in events:
        ws.append([ev.get(k) for k in keys])
    wb.save(path)
    return path
